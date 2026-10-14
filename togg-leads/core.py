"""Gemeinsame Bausteine: CSV-Tabellen, Einstellungen, Login/CSRF und Helfer."""

import csv
import io
import json
import os
import re
import secrets
import threading
import uuid
from contextlib import contextmanager
from datetime import datetime
from functools import wraps
from pathlib import Path

from flask import Response, abort, redirect, request, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("TOGG_DATA_DIR", BASE_DIR / "data"))
MEDIA_DIR = DATA_DIR / "media"

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"^\+?[0-9 ()/-]{6,20}$")
ZIP_RE = re.compile(r"^[0-9]{4,5}$")

TS_FORMAT = "%Y-%m-%d %H:%M:%S"


def now():
    return datetime.now().strftime(TS_FORMAT)


def parse_ts(value):
    try:
        return datetime.strptime(value, TS_FORMAT)
    except (TypeError, ValueError):
        return None


def safe_cell(value):
    """Schützt vor Formel-Injection, wenn ein Export in Excel geöffnet wird."""
    if PHONE_RE.match(value or ""):  # geprüfte Telefonnummer wie „+49 …“ unverändert lassen
        return value
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value


# --------------------------------------------------------------------------- #
# CSV-Tabellen
# --------------------------------------------------------------------------- #
class Table:
    """Eine CSV-Datei als einfache Tabelle. Alle Zugriffe laufen über ein Lock."""

    def __init__(self, filename, fields):
        self.path = DATA_DIR / filename
        self.fields = fields
        self._lock = threading.RLock()

    def _ensure(self):
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        if not self.path.exists():
            with self.path.open("w", newline="", encoding="utf-8") as f:
                csv.DictWriter(f, fieldnames=self.fields).writeheader()
            return
        with self.path.open(newline="", encoding="utf-8") as f:
            header = next(csv.reader(f), [])
        if header != self.fields:  # ältere Datei mit anderen Spalten → umschreiben
            with self.path.open(newline="", encoding="utf-8") as f:
                rows = list(csv.DictReader(f))
            self._write([{k: (r.get(k) or "") for k in self.fields} for r in rows])

    def _read(self):
        self._ensure()
        with self.path.open(newline="", encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
        for row in rows:  # neue Spalten in alten Dateien auffüllen
            for field in self.fields:
                row.setdefault(field, "")
                if row[field] is None:
                    row[field] = ""
        return rows

    def _write(self, rows):
        tmp = self.path.with_suffix(".tmp")
        with tmp.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=self.fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        tmp.replace(self.path)

    def all(self):
        with self._lock:
            return self._read()

    def append(self, row):
        with self._lock:
            self._ensure()
            with self.path.open("a", newline="", encoding="utf-8") as f:
                csv.DictWriter(f, fieldnames=self.fields, extrasaction="ignore").writerow(
                    {k: row.get(k, "") for k in self.fields}
                )

    @contextmanager
    def edit(self):
        """Lesen, ändern, zurückschreiben – atomar gegenüber anderen Anfragen."""
        with self._lock:
            rows = self._read()
            yield rows
            self._write(rows)


leads_table = Table(
    "leads.csv",
    ["id", "created_at", "name", "email", "phone", "zip_code", "interests", "notes", "consent", "source"],
)
queue_table = Table(
    "queue.csv",
    [
        "id", "token", "number", "day", "name", "phone", "email", "consent", "status",
        "created_at", "called_at", "started_at", "ended_at", "vehicle", "note",
    ],
)


# --------------------------------------------------------------------------- #
# Einstellungen (data/settings.json)
# --------------------------------------------------------------------------- #
DEFAULTS = {
    # Marke
    "brand_name": "Togg",
    "logo": "",
    "favicon": "",
    "accent_color": "#5aafdc",
    "public_url": "",
    # Module
    "module_leads": True,
    "module_queue": True,
    "tracking_enabled": False,
    "queue_also_lead": True,
    "show_eta": True,
    # Probefahrt
    "cars_count": 2,
    "drive_minutes": 20,
    "vehicles": "",
    # Formular
    "page_title": "Entdecken Sie {marke}.",
    "headline": "Entdecken Sie",
    "headline_accent": "{marke}.",
    "subline": "Sie interessieren sich für {marke}? Hinterlassen Sie uns gern Ihre Kontaktdaten.",
    "interests": "Testfahrt\nFinanzierung",
    "notes_placeholder": "Modell, Wunschtermin, Fragen …",
    "submit_label": "Interesse anmelden →",
    "consent_text": (
        "Ich willige ein, dass die {firma} die von mir angegebenen Daten speichert und mich zu meinem "
        "ausgewählten Interesse per E-Mail und Telefon kontaktiert. Diese Einwilligung ist freiwillig. "
        "Ich kann sie jederzeit mit Wirkung für die Zukunft widerrufen, z. B. per E-Mail an {email}. "
        "Die Rechtmäßigkeit der bis zum Widerruf erfolgten Verarbeitung bleibt davon unberührt."
    ),
    "thanks_title": "Vielen",
    "thanks_title_accent": "Dank.",
    "thanks_text": "Ihr Interesse wurde erfolgreich angemeldet. Wir melden uns in Kürze bei Ihnen.",
    # Warteschlange
    "queue_headline": "Probefahrt",
    "queue_headline_accent": "ohne Warten.",
    "queue_subline": (
        "Melden Sie sich an – wir rufen Sie auf, sobald ein Fahrzeug frei ist. "
        "Sie können sich in der Zwischenzeit frei umsehen."
    ),
    "queue_submit_label": "In Warteschlange eintragen →",
    "queue_consent_text": (
        "Ich willige ein, dass die {firma} meine Daten zur Organisation der Probefahrt speichert und "
        "mich dafür per Telefon, SMS oder WhatsApp kontaktiert. Widerruf jederzeit an {email}."
    ),
    "queue_waiting_text": "Lassen Sie diese Seite geöffnet – wir benachrichtigen Sie hier, sobald Sie dran sind.",
    "queue_called_text": "Bitte kommen Sie jetzt zum {marke}-Stand. Ihr Fahrzeug wartet!",
    "contact_message": "Hallo {name}, Sie sind jetzt dran für Ihre {marke}-Probefahrt (Nr. {nummer}). Bitte kommen Sie zum Stand.",
    # Rechtliches
    "company": "Togg Europe GmbH",
    "company_address": "Schelmenwasenstr. 34, 70567 Stuttgart",
    "privacy_email": "toggcare@togg.eu",
    "legal_text": (
        "Verantwortlich für die Datenverarbeitung ist die {firma}, {adresse}. Ihre Daten werden "
        "ausschließlich zur Bearbeitung Ihrer Anfrage und zur Kontaktaufnahme auf Grundlage Ihrer "
        "Einwilligung (Art. 6 Abs. 1 lit. a DSGVO) verarbeitet und nicht zu anderen Zwecken genutzt."
    ),
    "privacy_text": (
        "Verantwortlicher: {firma}, {adresse}\n"
        "Zweck: Bearbeitung Ihrer Anfrage bzw. Organisation Ihrer Probefahrt und Kontaktaufnahme\n"
        "Rechtsgrundlage: Einwilligung, Art. 6 Abs. 1 lit. a DSGVO\n"
        "Gespeicherte Daten: Name, Kontaktdaten, Interesse, Notizen, ggf. Zeitpunkte der Probefahrt\n"
        "Widerruf: jederzeit per E-Mail an {email}\n\n"
        "Hinweis: Platzhalter – bitte durch die vollständige Datenschutzerklärung ersetzen."
    ),
    # intern
    "admin_password_hash": "",
}

# (Abschnitt, [(Schlüssel, Bezeichnung, Typ, Hilfetext)])
SETTINGS_SCHEMA = [
    ("Marke", [
        ("brand_name", "Markenname", "text", "Ersetzt {marke} in allen Texten."),
        ("logo", "Logo", "file", "PNG, SVG, JPG oder WebP, max. 2 MB. Am besten mit transparentem Hintergrund."),
        ("favicon", "Favicon (Browser-Tab)", "file", "Quadratisch, z. B. 64×64 PNG."),
        ("accent_color", "Akzentfarbe", "color", "Buttons, Hervorhebungen, Auswahl."),
        ("public_url", "Öffentliche Adresse (optional)", "text",
         "Für QR-Codes, z. B. https://mein-pc.tail1234.ts.net – leer = aktuelle Adresse."),
    ]),
    ("Module", [
        ("module_leads", "Interessenten-Formular aktiv", "bool", "Startseite /"),
        ("module_queue", "Probefahrt-Warteschlange aktiv", "bool", "Seite /warteschlange"),
        ("tracking_enabled", "Fahrtzeit-Tracking", "bool",
         "Start und Ende jeder Probefahrt erfassen (Dauer, Fahrzeug, Durchschnitt)."),
        ("queue_also_lead", "Warteschlangen-Anmeldungen auch als Interessent speichern", "bool", ""),
        ("show_eta", "Geschätzte Wartezeit anzeigen", "bool", ""),
    ]),
    ("Probefahrt", [
        ("cars_count", "Anzahl Testfahrzeuge", "int", "Für die Wartezeit-Schätzung."),
        ("drive_minutes", "Ø Dauer einer Probefahrt (Minuten)", "int",
         "Mit Tracking wird automatisch der echte Durchschnitt genutzt."),
        ("vehicles", "Fahrzeuge (eins pro Zeile, optional)", "lines", "z. B. „T10X Weiß · S-TG 101“"),
        ("contact_message", "Nachricht für SMS/WhatsApp", "textarea", "Platzhalter: {name}, {nummer}, {marke}"),
    ]),
    ("Texte – Interessenten-Formular", [
        ("page_title", "Seitentitel (Browser-Tab)", "text", ""),
        ("headline", "Überschrift", "text", ""),
        ("headline_accent", "Überschrift – farbiger Teil", "text", ""),
        ("subline", "Untertitel", "textarea", ""),
        ("interests", "Interessen (eins pro Zeile)", "lines", ""),
        ("notes_placeholder", "Platzhalter Notizen", "text", ""),
        ("submit_label", "Button-Text", "text", ""),
        ("consent_text", "Einwilligungstext", "textarea", ""),
        ("thanks_title", "Danke-Seite: Überschrift", "text", ""),
        ("thanks_title_accent", "Danke-Seite: farbiger Teil", "text", ""),
        ("thanks_text", "Danke-Seite: Text", "textarea", ""),
    ]),
    ("Texte – Warteschlange", [
        ("queue_headline", "Überschrift", "text", ""),
        ("queue_headline_accent", "Überschrift – farbiger Teil", "text", ""),
        ("queue_subline", "Untertitel", "textarea", ""),
        ("queue_submit_label", "Button-Text", "text", ""),
        ("queue_consent_text", "Einwilligungstext", "textarea", ""),
        ("queue_waiting_text", "Hinweis während des Wartens", "textarea", ""),
        ("queue_called_text", "Text, wenn aufgerufen", "textarea", ""),
    ]),
    ("Rechtliches", [
        ("company", "Verantwortliche Firma", "text", "Ersetzt {firma}."),
        ("company_address", "Adresse", "text", "Ersetzt {adresse}."),
        ("privacy_email", "Datenschutz-E-Mail", "text", "Ersetzt {email}."),
        ("legal_text", "Hinweis unter dem Formular", "textarea", ""),
        ("privacy_text", "Datenschutzerklärung", "textarea", ""),
    ]),
]

SETTINGS_PATH = DATA_DIR / "settings.json"
_settings_lock = threading.Lock()


def get_settings():
    data = dict(DEFAULTS)
    if SETTINGS_PATH.exists():
        try:
            data.update(json.loads(SETTINGS_PATH.read_text(encoding="utf-8")))
        except (OSError, ValueError):
            pass
    return data


def save_settings(updates):
    with _settings_lock:
        data = get_settings()
        data.update(updates)
        stored = {k: v for k, v in data.items() if k in DEFAULTS}
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        tmp = SETTINGS_PATH.with_suffix(".tmp")
        tmp.write_text(json.dumps(stored, ensure_ascii=False, indent=2), encoding="utf-8")
        tmp.replace(SETTINGS_PATH)


def lines(value):
    return [line.strip() for line in str(value).splitlines() if line.strip()]


def fill(text, cfg, **extra):
    """Ersetzt {marke}, {firma}, {adresse}, {email} (und Extras) in einem Text."""
    values = {
        "marke": cfg["brand_name"],
        "firma": cfg["company"],
        "adresse": cfg["company_address"],
        "email": cfg["privacy_email"],
        **extra,
    }
    text = str(text)
    for key, value in values.items():
        text = text.replace("{" + key + "}", str(value))
    return text


# --------------------------------------------------------------------------- #
# Login & CSRF
# --------------------------------------------------------------------------- #
ENV_ADMIN_PASSWORD = os.environ.get("TOGG_ADMIN_PASSWORD", "admin")


def check_admin_password(password):
    stored = get_settings().get("admin_password_hash")
    if stored:
        return check_password_hash(stored, password)
    return secrets.compare_digest(password, ENV_ADMIN_PASSWORD)


def set_admin_password(password):
    save_settings({"admin_password_hash": generate_password_hash(password)})


def using_default_password():
    return not get_settings().get("admin_password_hash") and ENV_ADMIN_PASSWORD == "admin"


def login_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not session.get("admin"):
            return redirect(url_for("admin_login", next=request.full_path))
        return view(*args, **kwargs)

    return wrapper


def csrf_token():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_hex(16)
    return session["csrf_token"]


def check_csrf():
    token = request.form.get("csrf_token", "") or request.headers.get("X-CSRF-Token", "")
    if not token or not secrets.compare_digest(token, session.get("csrf_token", "")):
        abort(400)


# --------------------------------------------------------------------------- #
# Interessenten & Exporte
# --------------------------------------------------------------------------- #
def de_date(value):
    ts = parse_ts(value)
    return ts.strftime("%d.%m.%Y, %H:%M") if ts else value


def add_lead(name, email, phone, zip_code, interests, notes, source="Formular"):
    leads_table.append(
        {
            "id": uuid.uuid4().hex[:10],
            "created_at": now(),
            "name": name,
            "email": email,
            "phone": phone,
            "zip_code": zip_code,
            "interests": "|".join(interests),
            "notes": notes,
            "consent": now(),
            "source": source,
        }
    )


def csv_response(header, rows, name):
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(header)
    writer.writerows(rows)
    return Response(
        "\ufeff" + buf.getvalue(),  # BOM, damit Excel Umlaute korrekt anzeigt
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={name}-{datetime.now():%Y-%m-%d}.csv"},
    )


def xlsx_response(header, rows, name, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    color = get_settings()["accent_color"].lstrip("#").upper()
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color if len(color) == 6 else "5AAFDC")
    for row in rows:
        ws.append(row)
    for idx, col in enumerate(zip(*([header] + rows)), start=1):
        width = min(max(len(str(v)) for v in col) + 2, 50)
        ws.column_dimensions[get_column_letter(idx)].width = max(width, 8)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={name}-{datetime.now():%Y-%m-%d}.xlsx"},
    )
