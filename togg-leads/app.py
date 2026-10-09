"""Togg Leads – Interessenten-Formular mit lokaler CSV-Datenbank und Admin-Ansicht.

Öffentlich:  /          Formular „Entdecken Sie Togg.“
Admin:       /admin     Interessenten-Übersicht (Passwort-geschützt)
"""

import csv
import io
import os
import re
import secrets
import threading
import uuid
from datetime import datetime
from functools import wraps
from pathlib import Path

from flask import (
    Flask,
    Response,
    abort,
    flash,
    redirect,
    render_template,
    request,
    session,
    url_for,
)

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("TOGG_DATA_DIR", BASE_DIR / "data"))
CSV_PATH = DATA_DIR / "leads.csv"
SECRET_PATH = DATA_DIR / ".secret_key"

ADMIN_PASSWORD = os.environ.get("TOGG_ADMIN_PASSWORD", "admin")

FIELDS = ["id", "created_at", "name", "email", "phone", "zip_code", "interests", "notes", "consent"]
INTERESTS = ["Testfahrt", "Finanzierung"]

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"^\+?[0-9 ()/-]{6,20}$")
ZIP_RE = re.compile(r"^[0-9]{4,5}$")

_lock = threading.Lock()


def _secret_key():
    if os.environ.get("TOGG_SECRET_KEY"):
        return os.environ["TOGG_SECRET_KEY"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not SECRET_PATH.exists():
        SECRET_PATH.write_text(secrets.token_hex(32))
    return SECRET_PATH.read_text().strip()


app = Flask(__name__)
app.secret_key = _secret_key()


# --------------------------------------------------------------------------- #
# CSV-"Datenbank"
# --------------------------------------------------------------------------- #
def _ensure_csv():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not CSV_PATH.exists():
        with CSV_PATH.open("w", newline="", encoding="utf-8") as f:
            csv.DictWriter(f, fieldnames=FIELDS).writeheader()


def load_leads():
    _ensure_csv()
    with _lock, CSV_PATH.open(newline="", encoding="utf-8") as f:
        leads = list(csv.DictReader(f))
    for lead in leads:
        lead["interest_list"] = [i for i in lead.get("interests", "").split("|") if i]
    return leads


def save_leads(leads):
    _ensure_csv()
    tmp = CSV_PATH.with_suffix(".tmp")
    with _lock:
        with tmp.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=FIELDS, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(leads)
        tmp.replace(CSV_PATH)


def add_lead(lead):
    _ensure_csv()
    with _lock, CSV_PATH.open("a", newline="", encoding="utf-8") as f:
        csv.DictWriter(f, fieldnames=FIELDS).writerow(lead)


def _safe(value):
    """Schützt vor Formel-Injection, wenn der Export in Excel geöffnet wird."""
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value


@app.template_filter("de_date")
def de_date(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d %H:%M:%S").strftime("%d.%m.%Y, %H:%M")
    except ValueError:
        return value


# --------------------------------------------------------------------------- #
# CSRF
# --------------------------------------------------------------------------- #
@app.context_processor
def inject_csrf():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_hex(16)
    return {"csrf_token": session["csrf_token"]}


def _check_csrf():
    token = request.form.get("csrf_token", "")
    if not token or not secrets.compare_digest(token, session.get("csrf_token", "")):
        abort(400)


# --------------------------------------------------------------------------- #
# Öffentliches Formular
# --------------------------------------------------------------------------- #
def validate(form):
    data = {
        "name": form.get("name", "").strip(),
        "email": form.get("email", "").strip(),
        "phone": form.get("phone", "").strip(),
        "zip_code": form.get("zip_code", "").strip(),
        "notes": form.get("notes", "").strip(),
        "interests": [i for i in form.getlist("interests") if i in INTERESTS],
        "consent": bool(form.get("consent")),
    }
    errors = {}
    if not data["name"]:
        errors["name"] = "Bitte geben Sie Ihren Namen an."
    if not EMAIL_RE.match(data["email"]):
        errors["email"] = "Bitte geben Sie eine gültige E-Mail-Adresse an."
    if data["phone"] and not PHONE_RE.match(data["phone"]):
        errors["phone"] = "Bitte geben Sie eine gültige Telefonnummer an."
    if data["zip_code"] and not ZIP_RE.match(data["zip_code"]):
        errors["zip_code"] = "Bitte geben Sie eine gültige Postleitzahl an."
    if not data["interests"]:
        errors["interests"] = "Bitte wählen Sie mindestens ein Interesse."
    if len(data["notes"]) > 2000:
        errors["notes"] = "Maximal 2000 Zeichen."
    if not data["consent"]:
        errors["consent"] = "Bitte bestätigen Sie die Einwilligung."
    return data, errors


@app.route("/", methods=["GET", "POST"])
def index():
    data, errors = {"interests": []}, {}
    if request.method == "POST":
        _check_csrf()
        if request.form.get("website"):  # Honeypot gegen Spam-Bots
            return redirect(url_for("thanks"))
        data, errors = validate(request.form)
        if not errors:
            add_lead(
                {
                    "id": uuid.uuid4().hex[:10],
                    "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "name": data["name"],
                    "email": data["email"],
                    "phone": data["phone"],
                    "zip_code": data["zip_code"],
                    "interests": "|".join(data["interests"]),
                    "notes": data["notes"],
                    "consent": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
            )
            return redirect(url_for("thanks"))
    return render_template("index.html", data=data, errors=errors, interests=INTERESTS)


@app.route("/danke")
def thanks():
    return render_template("thanks.html")


@app.route("/datenschutz")
def privacy():
    return render_template("privacy.html")


# --------------------------------------------------------------------------- #
# Admin
# --------------------------------------------------------------------------- #
def login_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not session.get("admin"):
            return redirect(url_for("admin_login"))
        return view(*args, **kwargs)

    return wrapper


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        _check_csrf()
        if secrets.compare_digest(request.form.get("password", ""), ADMIN_PASSWORD):
            session["admin"] = True
            return redirect(url_for("admin"))
        flash("Falsches Passwort.")
    return render_template("login.html")


@app.route("/admin/logout", methods=["POST"])
def admin_logout():
    _check_csrf()
    session.pop("admin", None)
    return redirect(url_for("admin_login"))


def _filtered(leads, q, interest):
    q = q.lower()
    result = [
        l
        for l in reversed(leads)  # neueste zuerst, auch bei gleicher Sekunde
        if (not interest or interest in l["interest_list"])
        and (not q or q in " ".join([l["name"], l["email"], l["phone"], l["zip_code"], l["notes"]]).lower())
    ]
    result.sort(key=lambda l: l["created_at"], reverse=True)
    return result


@app.route("/admin")
@login_required
def admin():
    leads = load_leads()
    q = request.args.get("q", "").strip()
    interest = request.args.get("interest", "")
    if interest not in INTERESTS:
        interest = ""
    today = datetime.now().strftime("%Y-%m-%d")
    stats = {
        "total": len(leads),
        "Testfahrt": sum("Testfahrt" in l["interest_list"] for l in leads),
        "Finanzierung": sum("Finanzierung" in l["interest_list"] for l in leads),
        "today": sum(l["created_at"].startswith(today) for l in leads),
    }
    return render_template(
        "admin.html",
        leads=_filtered(leads, q, interest),
        stats=stats,
        q=q,
        interest=interest,
        interests=INTERESTS,
    )


@app.route("/admin/lead/<lead_id>/delete", methods=["POST"])
@login_required
def admin_delete(lead_id):
    _check_csrf()
    leads = load_leads()
    remaining = [l for l in leads if l["id"] != lead_id]
    if len(remaining) == len(leads):
        abort(404)
    save_leads(remaining)
    return redirect(url_for("admin", q=request.form.get("q", ""), interest=request.form.get("interest", "")))


EXPORT_HEADER = ["Datum", "Name", "E-Mail", "Telefon", "PLZ", "Testfahrt", "Finanzierung", "Notizen", "Einwilligung"]


def _export_rows():
    leads = _filtered(load_leads(), request.args.get("q", ""), request.args.get("interest", ""))
    for l in leads:
        yield [
            de_date(l["created_at"]),
            _safe(l["name"]),
            _safe(l["email"]),
            _safe(l["phone"]),
            l["zip_code"],
            "Ja" if "Testfahrt" in l["interest_list"] else "Nein",
            "Ja" if "Finanzierung" in l["interest_list"] else "Nein",
            _safe(l["notes"]),
            de_date(l["consent"]),
        ]


def _filename(ext):
    return f"togg-interessenten-{datetime.now():%Y-%m-%d}.{ext}"


@app.route("/admin/export.csv")
@login_required
def export_csv():
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(EXPORT_HEADER)
    writer.writerows(_export_rows())
    return Response(
        "﻿" + buf.getvalue(),  # BOM, damit Excel Umlaute korrekt anzeigt
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={_filename('csv')}"},
    )


@app.route("/admin/export.xlsx")
@login_required
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Interessenten"
    ws.append(EXPORT_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5AAFDC")
    for row in _export_rows():
        ws.append(row)
    for col, width in zip("ABCDEFGHI", [18, 24, 30, 18, 8, 11, 13, 40, 18]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={_filename('xlsx')}"},
    )


if __name__ == "__main__":
    _ensure_csv()
    if ADMIN_PASSWORD == "admin":
        print("WARNUNG: Standard-Admin-Passwort 'admin' aktiv – setze TOGG_ADMIN_PASSWORD!")
    app.run(
        host=os.environ.get("HOST", "127.0.0.1"),
        port=int(os.environ.get("PORT", 5000)),
        debug=False,
    )
